import os
import subprocess
import time
import json
from datetime import datetime
import re
from flask import Flask, request, render_template, send_file, jsonify, session
from werkzeug.utils import secure_filename
from docx import Document
from docx.shared import Inches, Pt, RGBColor
import PyPDF2
import requests
from dotenv import load_dotenv
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from io import BytesIO

# Load environment variables
load_dotenv()

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['RESULTS_FOLDER'] = 'results'
app.config['MAX_CONTENT_LENGTH'] = 32 * 1024 * 1024  # 32MB max file size
app.secret_key = os.urandom(24)  # For session management

# Create directories if they don't exist
for directory in [app.config['UPLOAD_FOLDER'], app.config['RESULTS_FOLDER']]:
    if not os.path.exists(directory):
        os.makedirs(directory)

# Store statistics
STATS_FILE = 'results/usage_statistics.json'
def update_statistics(file_type, question_count, generation_time):
    stats = {}
    if os.path.exists(STATS_FILE):
        with open(STATS_FILE, 'r') as f:
            stats = json.load(f)
    
    date = datetime.now().strftime('%Y-%m-%d')
    if date not in stats:
        stats[date] = {'total_files': 0, 'file_types': {}, 'total_questions': 0, 'avg_time': 0}
    
    stats[date]['total_files'] += 1
    stats[date]['file_types'][file_type] = stats[date]['file_types'].get(file_type, 0) + 1
    stats[date]['total_questions'] += question_count
    stats[date]['avg_time'] = ((stats[date]['avg_time'] * (stats[date]['total_files'] - 1)) + generation_time) / stats[date]['total_files']
    
    with open(STATS_FILE, 'w') as f:
        json.dump(stats, f)

# Question difficulty estimator
def estimate_difficulty(question):
    # Keywords indicating difficulty
    easy_keywords = ['basic', 'simple', 'straightforward', 'define', 'list', 'name']
    medium_keywords = ['explain', 'describe', 'compare', 'analyze', 'discuss']
    hard_keywords = ['evaluate', 'synthesize', 'critique', 'hypothesize', 'predict']
    
    question_lower = question.lower()
    
    if any(keyword in question_lower for keyword in hard_keywords):
        return 'Hard'
    elif any(keyword in question_lower for keyword in medium_keywords):
        return 'Medium'
    else:
        return 'Easy'

# Ollama configuration
OLLAMA_API_URL = "http://localhost:11434/api/generate"
MODEL_NAME = "mistral"  # or any other model you prefer

def ensure_ollama_running():
    # Check if Ollama is running
    try:
        response = requests.get("http://localhost:11434/api/tags")
        if response.status_code == 200:
            return True
    except requests.exceptions.ConnectionError:
        pass
    
    # Start Ollama if it's not running
    try:
        if os.name == 'nt':  # Windows
            subprocess.Popen(["ollama", "serve"], 
                           creationflags=subprocess.CREATE_NEW_CONSOLE)
        else:  # Unix/Linux/MacOS
            subprocess.Popen(["ollama", "serve"], 
                           start_new_session=True)
        
        # Wait for Ollama to start
        for _ in range(30):  # Try for 30 seconds
            try:
                response = requests.get("http://localhost:11434/api/tags")
                if response.status_code == 200:
                    # Pull the model if not already present
                    subprocess.run(["ollama", "pull", MODEL_NAME])
                    return True
            except requests.exceptions.ConnectionError:
                time.sleep(1)
                continue
    except Exception as e:
        print(f"Error starting Ollama: {e}")
        return False
    
    return False

# Ensure Ollama is running when the app starts
if not ensure_ollama_running():
    raise RuntimeError("Failed to start Ollama service")

ALLOWED_EXTENSIONS = {'txt', 'pdf', 'docx'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_from_pdf(file_path):
    text = ""
    with open(file_path, 'rb') as file:
        pdf_reader = PyPDF2.PdfReader(file)
        for page in pdf_reader.pages:
            text += page.extract_text()
    return text

def extract_text_from_docx(file_path):
    doc = Document(file_path)
    return " ".join([paragraph.text for paragraph in doc.paragraphs])

def generate_mcqs(text, num_questions=5, difficulty=None, topic=None):
    # Prepare custom prompt based on parameters
    prompt_template = f"""Generate {num_questions} multiple choice questions from the following text.
    Requirements:
    1. Each question should have one correct answer and three incorrect answers
    2. Questions should be {difficulty if difficulty else 'mixed'} difficulty
    3. Focus on {topic if topic else 'key concepts'}
    4. Format each question as follows:
       Q[number]. [Question text]
       A) [Option]
       B) [Option]
       C) [Option]
       D) [Option]
       Correct Answer: [A/B/C/D]
       Explanation: [Brief explanation why this is correct]
       Difficulty: [Easy/Medium/Hard]
       Topic: [Subject area]
    
    Text: {text}"""
    
    start_time = time.time()
    
    response = requests.post(OLLAMA_API_URL, json={
        "model": MODEL_NAME,
        "prompt": prompt_template,
        "stream": False
    })
    
    if response.status_code != 200:
        raise Exception("Failed to generate MCQs")
    
    generation_time = time.time() - start_time
    mcq_text = response.json()['response']
    
    # Parse and structure the MCQs
    questions = []
    current_question = {}
    
    for line in mcq_text.split('\n'):
        line = line.strip()
        if not line:
            continue
        
        if line.startswith('Q'):
            if current_question:
                questions.append(current_question)
            current_question = {'options': []}
            current_question['question'] = line[line.find('.')+1:].strip()
        elif line.startswith(('A)', 'B)', 'C)', 'D)')):
            current_question['options'].append(line[2:].strip())
        elif line.startswith('Correct Answer:'):
            current_question['correct_answer'] = line.split(':')[1].strip()
        elif line.startswith('Explanation:'):
            current_question['explanation'] = line.split(':')[1].strip()
        elif line.startswith('Difficulty:'):
            current_question['difficulty'] = line.split(':')[1].strip()
        elif line.startswith('Topic:'):
            current_question['topic'] = line.split(':')[1].strip()
    
    if current_question:
        questions.append(current_question)
    
    # Store the structured MCQs in session
    session['last_mcqs'] = questions
    session['generation_time'] = generation_time
    
    return format_mcqs(questions), len(questions), generation_time

def format_mcqs(questions):
    formatted_text = "MCQ Generator Results\n"
    formatted_text += "=" * 50 + "\n\n"
    
    for i, q in enumerate(questions, 1):
        formatted_text += f"Question {i}: {q['question']}\n"
        for j, opt in enumerate(q['options']):
            formatted_text += f"{chr(65+j)}) {opt}\n"
        formatted_text += f"\nCorrect Answer: {q['correct_answer']}\n"
        formatted_text += f"Explanation: {q['explanation']}\n"
        formatted_text += f"Difficulty: {q['difficulty']}\n"
        formatted_text += f"Topic: {q['topic']}\n"
        formatted_text += "=" * 50 + "\n\n"
    
    return formatted_text

@app.route('/', methods=['GET'])
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({'error': 'No file part'}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No selected file'}), 400
    
    if file and allowed_file(file.filename):
        filename = secure_filename(file.filename)
        file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(file_path)
        
        # Get parameters from request
        num_questions = int(request.form.get('num_questions', 5))
        difficulty = request.form.get('difficulty', None)
        topic = request.form.get('topic', None)
        
        # Extract text based on file type
        file_type = filename.split('.')[-1].lower()
        if file_type == 'pdf':
            text = extract_text_from_pdf(file_path)
        elif file_type == 'docx':
            text = extract_text_from_docx(file_path)
        else:  # txt file
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()
        
        # Generate MCQs
        mcqs, question_count, generation_time = generate_mcqs(text, num_questions, difficulty, topic)
        
        # Update statistics
        update_statistics(file_type, question_count, generation_time)
        
        # Save results in different formats
        base_name = f'generated_mcqs_{datetime.now().strftime("%Y%m%d_%H%M%S")}'
        
        # Save as TXT
        txt_file = os.path.join(app.config['RESULTS_FOLDER'], f'{base_name}.txt')
        with open(txt_file, 'w', encoding='utf-8') as f:
            f.write(mcqs)
        
        # Save as DOCX with formatting
        docx_file = os.path.join(app.config['RESULTS_FOLDER'], f'{base_name}.docx')
        create_formatted_docx(session['last_mcqs'], docx_file)
        
        return jsonify({
            'message': 'MCQs generated successfully',
            'txt_file': f'{base_name}.txt',
            'docx_file': f'{base_name}.docx',
            'stats': {
                'question_count': question_count,
                'generation_time': round(generation_time, 2),
                'file_type': file_type
            }
        })
    
    return jsonify({'error': 'File type not allowed'}), 400

@app.route('/stats')
def get_statistics():
    if not os.path.exists(STATS_FILE):
        return jsonify({'error': 'No statistics available'}), 404
    
    with open(STATS_FILE, 'r') as f:
        stats = json.load(f)
    
    # Generate visualizations
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(15, 5))
    
    # File types pie chart
    file_types = {'pdf': 0, 'docx': 0, 'txt': 0}
    for date in stats:
        for file_type, count in stats[date]['file_types'].items():
            file_types[file_type] = file_types.get(file_type, 0) + count
    
    ax1.pie(file_types.values(), labels=file_types.keys(), autopct='%1.1f%%')
    ax1.set_title('File Types Distribution')
    
    # Daily questions line plot
    dates = list(stats.keys())
    questions = [stats[date]['total_questions'] for date in dates]
    ax2.plot(dates, questions, marker='o')
    ax2.set_title('Questions Generated per Day')
    ax2.tick_params(axis='x', rotation=45)
    
    # Save plot to buffer
    buf = BytesIO()
    plt.savefig(buf, format='png', bbox_inches='tight')
    buf.seek(0)
    plt.close()
    
    return jsonify({
        'stats': stats,
        'visualization': buf.getvalue().hex()
    })

@app.route('/download/<filename>')
def download_file(filename):
    return send_file(
        os.path.join(app.config['RESULTS_FOLDER'], filename),
        as_attachment=True
    )

def create_formatted_docx(questions, output_file):
    doc = Document()
    
    # Add title
    title = doc.add_heading('Generated Multiple Choice Questions', 0)
    title.alignment = 1  # Center alignment
    
    # Add timestamp
    doc.add_paragraph(f'Generated on: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
    doc.add_paragraph('=' * 50)
    
    # Add questions
    for i, q in enumerate(questions, 1):
        # Question
        question_para = doc.add_paragraph()
        question_run = question_para.add_run(f'Question {i}: {q["question"]}')
        question_run.bold = True
        
        # Options
        for j, opt in enumerate(q['options']):
            option_para = doc.add_paragraph()
            option_para.add_run(f'{chr(65+j)}) {opt}')
        
        # Correct Answer
        answer_para = doc.add_paragraph()
        answer_run = answer_para.add_run(f'Correct Answer: {q["correct_answer"]}')
        answer_run.bold = True
        answer_run.font.color.rgb = RGBColor(0, 128, 0)  # Green color
        
        # Explanation
        explanation_para = doc.add_paragraph()
        explanation_para.add_run('Explanation: ').bold = True
        explanation_para.add_run(q['explanation'])
        
        # Metadata
        meta_para = doc.add_paragraph()
        meta_para.add_run(f'Difficulty: {q["difficulty"]} | Topic: {q["topic"]}')
        meta_para.style = 'Subtle Emphasis'
        
        doc.add_paragraph('=' * 50)
    
    doc.save(output_file)

@app.route('/dashboard')
def dashboard():
    return render_template('dashboard.html')

@app.route('/api/templates', methods=['GET'])
def get_templates():
    templates = {
        'General': {
            'name': 'General Knowledge',
            'description': 'Basic multiple choice format',
            'structure': {
                'question_count': 5,
                'difficulty': 'mixed',
                'include_explanations': True,
                'format': 'standard'
            }
        },
        'Technical': {
            'name': 'Technical Questions',
            'description': 'Code-focused questions',
            'structure': {
                'question_count': 10,
                'difficulty': 'hard',
                'include_explanations': True,
                'format': 'code'
            }
        },
        'Academic': {
            'name': 'Academic Test',
            'description': 'Formal academic format',
            'structure': {
                'question_count': 15,
                'difficulty': 'medium',
                'include_explanations': False,
                'format': 'academic'
            }
        }
    }
    return jsonify(templates)

@app.route('/api/question-bank', methods=['GET'])
def get_question_bank():
    # Load questions from storage
    questions_file = os.path.join(app.config['RESULTS_FOLDER'], 'question_bank.json')
    if os.path.exists(questions_file):
        with open(questions_file, 'r', encoding='utf-8') as f:
            return jsonify(json.load(f))
    return jsonify([])

@app.route('/api/question-bank', methods=['POST'])
def save_to_question_bank():
    questions = request.json
    questions_file = os.path.join(app.config['RESULTS_FOLDER'], 'question_bank.json')
    
    existing_questions = []
    if os.path.exists(questions_file):
        with open(questions_file, 'r', encoding='utf-8') as f:
            existing_questions = json.load(f)
    
    # Add new questions
    existing_questions.extend(questions)
    
    with open(questions_file, 'w', encoding='utf-8') as f:
        json.dump(existing_questions, f, indent=2)
    
    return jsonify({'message': 'Questions saved successfully'})

@app.route('/download/<filename>')
def download_file(filename):
    return send_file(
        os.path.join(app.config['RESULTS_FOLDER'], filename),
        as_attachment=True
    )

@app.route('/api/settings', methods=['GET', 'POST'])
def handle_settings():
    settings_file = os.path.join(app.config['RESULTS_FOLDER'], 'settings.json')
    
    if request.method == 'POST':
        settings = request.json
        with open(settings_file, 'w', encoding='utf-8') as f:
            json.dump(settings, f, indent=2)
        return jsonify({'message': 'Settings saved successfully'})
    
    if os.path.exists(settings_file):
        with open(settings_file, 'r', encoding='utf-8') as f:
            return jsonify(json.load(f))
    
    # Default settings
    return jsonify({
        'default_question_count': 5,
        'default_difficulty': 'mixed',
        'default_format': 'txt',
        'include_explanations': True,
        'save_to_bank': True
    })

if __name__ == '__main__':
    # Create required directories if they don't exist
    os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
    os.makedirs(app.config['RESULTS_FOLDER'], exist_ok=True)
    
    # Initialize settings and question bank if they don't exist
    settings_file = os.path.join(app.config['RESULTS_FOLDER'], 'settings.json')
    if not os.path.exists(settings_file):
        with open(settings_file, 'w', encoding='utf-8') as f:
            json.dump({
                'default_question_count': 5,
                'default_difficulty': 'mixed',
                'default_format': 'txt',
                'include_explanations': True,
                'save_to_bank': True
            }, f, indent=2)
    
    questions_file = os.path.join(app.config['RESULTS_FOLDER'], 'question_bank.json')
    if not os.path.exists(questions_file):
        with open(questions_file, 'w', encoding='utf-8') as f:
            json.dump([], f)
    
    app.run(debug=True)